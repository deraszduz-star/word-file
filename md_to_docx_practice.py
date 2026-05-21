# -*- coding: utf-8 -*-
"""
Конвертер markdown-отчёта о преддипломной практике в Word-документ (.docx)
с форматированием согласно требованиям оформления (ГОСТ, методичка КФУ).
"""

import re
from docx import Document
from docx.shared import Pt, Cm, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml


def setup_document():
    """Создаёт документ с правильными полями и настройками."""
    doc = Document()
    section = doc.sections[0]
    section.left_margin = Mm(30)
    section.right_margin = Mm(10)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.page_width = Mm(210)
    section.page_height = Mm(297)

    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(14)
    font.color.rgb = RGBColor(0, 0, 0)
    pf = style.paragraph_format
    pf.line_spacing = 1.5
    pf.first_line_indent = Cm(1.25)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = parse_xml(
            f'<w:rFonts {nsdecls("w")} w:ascii="Times New Roman" '
            f'w:hAnsi="Times New Roman" w:cs="Times New Roman"/>'
        )
        rPr.append(rFonts)
    else:
        rFonts.set(qn('w:ascii'), 'Times New Roman')
        rFonts.set(qn('w:hAnsi'), 'Times New Roman')
        rFonts.set(qn('w:cs'), 'Times New Roman')

    return doc


def add_page_numbers(doc):
    """Добавляет нумерацию страниц внизу по центру, 12 пт."""
    section = doc.sections[0]
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = p.add_run()
    run.font.size = Pt(12)
    run.font.name = 'Times New Roman'
    fldChar1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
    run._element.append(fldChar1)

    run2 = p.add_run()
    run2.font.size = Pt(12)
    run2.font.name = 'Times New Roman'
    instrText = parse_xml(
        f'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>'
    )
    run2._element.append(instrText)

    run3 = p.add_run()
    run3.font.size = Pt(12)
    run3.font.name = 'Times New Roman'
    fldChar2 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
    run3._element.append(fldChar2)


def _apply_run_formatting(run, bold=False, italic=False, font_size=14):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(font_size)
    run.font.name = 'Times New Roman'


def add_paragraph(doc, text, align='justify', bold=False, italic=False,
                  font_size=14, first_indent=1.25,
                  space_before=0, space_after=0):
    """Добавляет параграф простым текстом (без markdown-инлайна)."""
    p = doc.add_paragraph()
    alignments = {
        'justify': WD_ALIGN_PARAGRAPH.JUSTIFY,
        'center': WD_ALIGN_PARAGRAPH.CENTER,
        'right': WD_ALIGN_PARAGRAPH.RIGHT,
        'left': WD_ALIGN_PARAGRAPH.LEFT,
    }
    p.alignment = alignments.get(align, WD_ALIGN_PARAGRAPH.JUSTIFY)
    p.paragraph_format.first_line_indent = Cm(first_indent)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        run = p.add_run(text)
        _apply_run_formatting(run, bold=bold, italic=italic, font_size=font_size)
    return p


def add_paragraph_inline(doc, text, align='justify', font_size=14,
                         first_indent=1.25):
    """Добавляет параграф с обработкой инлайн-markdown (**жирный**, *курсив*)."""
    p = doc.add_paragraph()
    alignments = {
        'justify': WD_ALIGN_PARAGRAPH.JUSTIFY,
        'center': WD_ALIGN_PARAGRAPH.CENTER,
        'right': WD_ALIGN_PARAGRAPH.RIGHT,
        'left': WD_ALIGN_PARAGRAPH.LEFT,
    }
    p.alignment = alignments.get(align, WD_ALIGN_PARAGRAPH.JUSTIFY)
    p.paragraph_format.first_line_indent = Cm(first_indent)
    p.paragraph_format.line_spacing = 1.5

    # Разбиваем по **жирным** и *курсивным* фрагментам
    pattern = re.compile(r'(\*\*[^*]+\*\*|\*[^*]+\*)')
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            run = p.add_run(text[pos:m.start()])
            _apply_run_formatting(run, font_size=font_size)
        token = m.group(0)
        if token.startswith('**'):
            run = p.add_run(token[2:-2])
            _apply_run_formatting(run, bold=True, font_size=font_size)
        else:
            run = p.add_run(token[1:-1])
            _apply_run_formatting(run, italic=True, font_size=font_size)
        pos = m.end()
    if pos < len(text):
        run = p.add_run(text[pos:])
        _apply_run_formatting(run, font_size=font_size)
    return p


def add_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = parse_xml(f'<w:tblPr {nsdecls("w")}/>')
        tbl.insert(0, tblPr)
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tblBorders>'
    )
    tblPr.append(borders)


def _strip_md(s):
    s = re.sub(r'\*\*(.+?)\*\*', r'\1', s)
    s = re.sub(r'\*(.+?)\*', r'\1', s)
    return s


def create_table(doc, headers, rows):
    num_cols = len(headers)
    table = doc.add_table(rows=1 + len(rows), cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    add_table_borders(table)

    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ''
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(_strip_md(header))
        run.bold = True
        run.font.size = Pt(12)
        run.font.name = 'Times New Roman'

    for row_idx, row_data in enumerate(rows):
        for col_idx, cell_text in enumerate(row_data):
            cell = table.rows[row_idx + 1].cells[col_idx]
            cell.text = ''
            p = cell.paragraphs[0]
            if col_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.0
            raw = str(cell_text)
            is_bold = raw.startswith('**') and raw.endswith('**')
            run = p.add_run(_strip_md(raw))
            run.bold = is_bold
            run.font.size = Pt(12)
            run.font.name = 'Times New Roman'
    return table


def parse_md_table(table_lines):
    headers = []
    rows = []
    for line in table_lines:
        line = line.strip()
        if not line.startswith('|'):
            continue
        cells = [c.strip() for c in line.split('|')[1:-1]]
        if all(re.match(r'^[-:]+$', c) for c in cells):
            continue
        if not headers:
            headers = cells
        else:
            rows.append(cells)
    return headers, rows


# --- Жёстко собранный титульный лист и содержание ---

def add_title_page(doc):
    """Титульный лист по образцу методички КФУ."""
    add_paragraph(doc, 'Министерство науки и высшего образования Российской Федерации',
                  align='center', bold=True, first_indent=0)
    add_paragraph(doc,
                  'Федеральное государственное автономное образовательное '
                  'учреждение высшего образования',
                  align='center', bold=True, first_indent=0)
    add_paragraph(doc, '«Казанский (Приволжский) федеральный университет»',
                  align='center', bold=True, first_indent=0)
    for _ in range(2):
        add_paragraph(doc, '', first_indent=0)
    add_paragraph(doc, 'Институт управления, экономики и финансов',
                  align='center', first_indent=0)
    add_paragraph(doc, 'Кафедра _______________________________',
                  align='center', first_indent=0)
    add_paragraph(doc, 'Направление подготовки: _______________________________',
                  align='center', first_indent=0)
    for _ in range(4):
        add_paragraph(doc, '', first_indent=0)
    add_paragraph(doc, 'ОТЧЁТ', align='center', bold=True, first_indent=0)
    add_paragraph(doc, 'о прохождении преддипломной практики',
                  align='center', bold=True, first_indent=0)
    for _ in range(2):
        add_paragraph(doc, '', first_indent=0)
    add_paragraph(doc, 'Тема выпускной квалификационной работы:',
                  align='center', bold=True, first_indent=0)
    add_paragraph(doc,
                  '«Реинжиниринг бизнес-процессов структурных подразделений '
                  'Казанского федерального университета»',
                  align='center', first_indent=0)
    add_paragraph(doc, '', first_indent=0)
    add_paragraph(doc, 'Место прохождения практики:',
                  align='center', bold=True, first_indent=0)
    add_paragraph(doc,
                  'Первичная профсоюзная организация работников КФУ '
                  'Общероссийского Профсоюза образования',
                  align='center', first_indent=0)
    add_paragraph(doc, '', first_indent=0)
    add_paragraph(doc, 'Сроки прохождения практики: с 08.05.2026 по 22.05.2026',
                  align='center', first_indent=0)
    for _ in range(3):
        add_paragraph(doc, '', first_indent=0)
    add_paragraph(doc, 'Выполнил(а) обучающийся(аяся): _______________________________',
                  align='left', first_indent=0)
    add_paragraph(doc, 'Руководитель практики от организации: _______________________________',
                  align='left', first_indent=0)
    add_paragraph(doc, 'Руководитель практики от КФУ: _______________________________',
                  align='left', first_indent=0)
    add_paragraph(doc, 'Оценка: _______________________________',
                  align='left', first_indent=0)
    for _ in range(4):
        add_paragraph(doc, '', first_indent=0)
    add_paragraph(doc, 'Казань — 2026', align='center', first_indent=0)
    doc.add_page_break()


def add_contents_page(doc):
    """Раздел СОДЕРЖАНИЕ."""
    add_paragraph(doc, 'СОДЕРЖАНИЕ', align='center', bold=True, first_indent=0)
    add_paragraph(doc, '', first_indent=0)
    items = [
        'Введение',
        'Раздел 1. Постановка цели и задач исследования',
        'Раздел 2. Анализ предметной области объекта исследования',
        '    2.1. Анализ текущего состояния бизнес-процессов подразделений (AS IS)',
        '    2.2. Проектирование целевых процессов (TO BE)',
        '    2.3. Оценка экономической эффективности реинжиниринга',
        'Раздел 3. Выводы по итогам проведённого исследования',
        'Заключение',
        'Список использованных источников',
        'Приложения',
    ]
    for it in items:
        add_paragraph(doc, it, align='left', first_indent=0)
    doc.add_page_break()


# --- Основной парсер тела отчёта ---

# Структурные секции верхнего уровня (## …) — каждая с новой страницы,
# заголовок по центру, заглавными, полужирно
TOP_SECTIONS = {
    'ВВЕДЕНИЕ',
    'ЗАКЛЮЧЕНИЕ',
    'СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ',
    'ПРИЛОЖЕНИЯ',
}


def process_markdown(doc, md_text):
    lines = md_text.split('\n')
    i = 0
    n = len(lines)

    # Пропускаем всё до конца содержания (мы сгенерировали титул и содержание сами)
    # Ищем строку "## ВВЕДЕНИЕ"
    while i < n:
        s = lines[i].strip()
        if s.startswith('## ВВЕДЕНИЕ'):
            break
        i += 1

    in_code_block = False

    while i < n:
        line = lines[i]
        stripped = line.strip()

        # Кодовые блоки игнорируем (в этом отчёте их нет, но на всякий случай)
        if stripped.startswith('```'):
            in_code_block = not in_code_block
            i += 1
            continue
        if in_code_block:
            i += 1
            continue

        # Горизонтальные разделители пропускаем
        if stripped == '---':
            i += 1
            continue

        # Пустые строки пропускаем
        if not stripped:
            i += 1
            continue

        # ## ЗАГОЛОВОК
        if stripped.startswith('## '):
            text = stripped[3:].strip()
            text_upper = text.upper()
            # Раздел верхнего уровня (РАЗДЕЛ 1/2/3 или ВВЕДЕНИЕ/ЗАКЛЮЧЕНИЕ/...)
            doc.add_page_break()
            add_paragraph(doc, text_upper, align='center', bold=True,
                          first_indent=0)
            # пустая строка после заголовка раздела
            add_paragraph(doc, '', first_indent=0)
            i += 1
            continue

        # ### подзаголовок
        if stripped.startswith('### '):
            text = stripped[4:].strip()
            # Подзаголовки в списке источников (Нормативные правовые акты и т.д.)
            # — оформляем по центру, полужирно, без новой страницы
            add_paragraph(doc, '', first_indent=0)
            add_paragraph(doc, text, align='center', bold=True, first_indent=0)
            add_paragraph(doc, '', first_indent=0)
            i += 1
            continue

        # Markdown-таблица
        if stripped.startswith('|'):
            table_lines = []
            while i < n and lines[i].strip().startswith('|'):
                table_lines.append(lines[i])
                i += 1
            headers, rows = parse_md_table(table_lines)
            if headers and rows:
                create_table(doc, headers, rows)
                add_paragraph(doc, '', first_indent=0)
            continue

        # Маркированный список (- ...)
        if stripped.startswith('- '):
            text = stripped[2:].strip()
            add_paragraph_inline(doc, '\u2013 ' + text, align='justify',
                                 first_indent=1.25)
            i += 1
            continue

        # Нумерованный список вида "1) ..." или "1. ..."
        if re.match(r'^\d+[\)\.]\s+', stripped):
            add_paragraph_inline(doc, stripped, align='justify',
                                 first_indent=1.25)
            i += 1
            continue

        # Обычный абзац (возможен инлайн-markdown)
        add_paragraph_inline(doc, stripped, align='justify', first_indent=1.25)
        i += 1


def main():
    src = '/projects/sandbox/word-file/Отчет_практика.md'
    dst = '/projects/sandbox/word-file/Отчет_практика.docx'

    with open(src, 'r', encoding='utf-8') as f:
        md_text = f.read()

    doc = setup_document()
    add_page_numbers(doc)

    add_title_page(doc)
    add_contents_page(doc)
    process_markdown(doc, md_text)

    doc.save(dst)
    print(f'Документ создан: {dst}')


if __name__ == '__main__':
    main()
